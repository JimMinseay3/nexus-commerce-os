import csv
import io
from decimal import Decimal, InvalidOperation

from django.db import transaction
from openpyxl import load_workbook

from core.models import (
    ChannelSKU, Company, ImportError, IntegrationConnection, IntegrationProvider, InventoryLedger,
    MappingVersion, Product, RawRecord, SKU, Store, Supplier, Warehouse,
)
from .datahub import payload_digest
from .inventory import move_inventory
from .mapping import apply_mapping, preview_mapping


RESOURCE_FIELDS = {
    "suppliers": ["code", "name", "contact_name", "email", "phone", "currency", "payment_terms_days", "default_lead_time_days"],
    "skus": ["spu", "product_name", "sku", "sku_name", "barcode", "supplier_code", "purchase_price", "currency", "moq", "case_pack", "safety_stock", "gross_weight_kg", "length_cm", "width_cm", "height_cm"],
    "channel_mappings": ["store_external_id", "internal_sku", "external_sku", "external_product_id", "fulfillment"],
    "opening_inventory": ["warehouse_code", "sku", "quantity", "unit_cost"],
}


def workbook_rows(file_obj, mapping=None, limit=None):
    if str(getattr(file_obj, "name", "")).lower().endswith(".csv"):
        content = file_obj.read()
        text = content.decode("utf-8-sig") if isinstance(content, bytes) else content
        reader = csv.reader(io.StringIO(text))
        headers = [str(x).strip() for x in next(reader)]
        iterator = reader
    else:
        workbook = load_workbook(file_obj, read_only=True, data_only=True)
        sheet = workbook.active
        iterator = sheet.iter_rows(values_only=True)
        headers = [str(x).strip() if x is not None else "" for x in next(iterator)]
    mapping = mapping or {header: header for header in headers}
    rows = []
    for values in iterator:
        source = {headers[index]: value for index, value in enumerate(values) if index < len(headers)}
        target = {target_name: source.get(source_name) for source_name, target_name in mapping.items() if target_name}
        if any(value not in (None, "") for value in target.values()):
            rows.append(target)
        if limit and len(rows) >= limit:
            break
    return headers, rows


def _flatten_canonical(resource_type, row):
    namespaces = {
        "suppliers": "supplier", "skus": "sku", "channel_mappings": "channel_mapping",
        "opening_inventory": "inventory",
    }
    namespace = namespaces.get(resource_type)
    if namespace and isinstance(row.get(namespace), dict):
        flattened = dict(row[namespace])
        if resource_type == "skus":
            if "code" in flattened:
                flattened["sku"] = flattened.pop("code")
            if "name" in flattened:
                flattened["sku_name"] = flattened.pop("name")
            flattened.update({f"product_{key}": value for key, value in row.get("product", {}).items()})
            if "product_spu" in flattened:
                flattened["spu"] = flattened.pop("product_spu")
            if row.get("supplier", {}).get("code"):
                flattened["supplier_code"] = row["supplier"]["code"]
        return flattened
    return row


def _rows_for_job(job, limit=None, actor=None, create_run=False):
    mapping_version_id = job.mapping.get("mapping_version_id") if isinstance(job.mapping, dict) else None
    if not mapping_version_id:
        headers, rows = workbook_rows(job.file, job.mapping, limit)
        return headers, rows, rows, None, []
    version = MappingVersion.objects.select_related("mapping_set").get(
        pk=mapping_version_id, mapping_set__company=job.company, mapping_set__resource_type=job.resource_type,
        status="published",
    )
    headers, source_rows = workbook_rows(job.file, None, limit)
    if create_run:
        run = preview_mapping(version, source_rows, actor)
        outputs, errors = run.output_sample, run.errors
    else:
        outputs, errors = [], []
        for index, source in enumerate(source_rows):
            output, row_errors = apply_mapping(source, version.rules, job.company)
            outputs.append(output)
            errors.extend({"row": index + 2, **error} for error in row_errors)
    rows = [_flatten_canonical(job.resource_type, row) for row in outputs]
    return headers, source_rows, rows, version, errors


def preview_import(job, limit=20):
    job.file.open("rb")
    try:
        headers, _, rows, version, errors = _rows_for_job(job, limit, job.created_by, create_run=True)
    finally:
        job.file.close()
    job.preview = rows
    job.status = "previewed"
    job.save(update_fields=["preview", "status", "updated_at"])
    return {"headers": headers, "required_fields": RESOURCE_FIELDS.get(job.resource_type, []), "rows": rows,
            "mapping_version": str(version.id) if version else None, "errors": errors}


@transaction.atomic
def execute_import(job, actor=None):
    job.errors.all().delete()
    job.file.open("rb")
    try:
        _, source_rows, rows, mapping_version, mapping_errors = _rows_for_job(job)
    finally:
        job.file.close()
    job.total_rows, success = len(rows), 0
    errors_by_row = {error["row"] for error in mapping_errors}
    provider = IntegrationProvider.objects.filter(key="file").first()
    connection = IntegrationConnection.objects.get_or_create(
        company=job.company, provider=provider, name="Excel / CSV 导入中心",
        defaults={"environment": "sandbox", "region": "LOCAL", "enabled_capabilities": provider.capabilities},
    )[0] if provider else None
    for index, row in enumerate(rows, start=2):
        source = source_rows[index - 2]
        if index in errors_by_row:
            row_errors = [error for error in mapping_errors if error["row"] == index]
            ImportError.objects.create(job=job, row_number=index, message=str(row_errors), raw_data=source)
            continue
        try:
            with transaction.atomic():
                entity = _import_row(job.company, job.resource_type, row, actor, index, job.id)
                if connection:
                    raw, _ = RawRecord.objects.get_or_create(
                        company=job.company, connection=connection, object_type=job.resource_type,
                        external_id=str(index), external_version=str(job.id), scope=getattr(job.file, "name", ""),
                        payload_hash=payload_digest(source),
                        defaults={"payload": source, "mapping_version": mapping_version, "status": "merged",
                                  "canonical_entity_type": entity._meta.model_name if entity else "",
                                  "canonical_entity_id": entity.id if entity else None},
                    )
            success += 1
        except Exception as exc:
            ImportError.objects.create(job=job, row_number=index, message=str(exc), raw_data=row)
    job.success_rows = success
    job.error_rows = len(rows) - success
    job.status = "completed" if success else "failed"
    job.save(update_fields=["total_rows", "success_rows", "error_rows", "status", "updated_at"])
    return job


def _decimal(value, default="0"):
    try:
        return Decimal(str(value if value not in (None, "") else default))
    except InvalidOperation as exc:
        raise ValueError(f"无效数字: {value}") from exc


def _import_row(company: Company, resource_type, row, actor, row_number, job_id):
    if resource_type == "suppliers":
        code, name = str(row.get("code") or "").strip(), str(row.get("name") or "").strip()
        if not code or not name:
            raise ValueError("供应商编码和名称必填")
        entity, _ = Supplier.objects.update_or_create(company=company, code=code, defaults={
            "name": name, "contact_name": row.get("contact_name") or "", "email": row.get("email") or "",
            "phone": row.get("phone") or "", "currency": row.get("currency") or "CNY",
            "payment_terms_days": int(row.get("payment_terms_days") or 0),
            "default_lead_time_days": int(row.get("default_lead_time_days") or 30),
        })
    elif resource_type == "skus":
        spu, sku_code = str(row.get("spu") or "").strip(), str(row.get("sku") or "").strip()
        if not spu or not sku_code:
            raise ValueError("SPU 和 SKU 必填")
        product, _ = Product.objects.update_or_create(company=company, spu=spu, defaults={"name": row.get("product_name") or spu, "status": Product.Status.ACTIVE})
        supplier = Supplier.objects.filter(company=company, code=row.get("supplier_code")).first()
        entity, _ = SKU.objects.update_or_create(company=company, code=sku_code, defaults={
            "product": product, "name": row.get("sku_name") or sku_code, "barcode": row.get("barcode") or "",
            "supplier": supplier, "purchase_price": _decimal(row.get("purchase_price")), "currency": row.get("currency") or "CNY",
            "moq": int(row.get("moq") or 1), "case_pack": int(row.get("case_pack") or 1), "safety_stock": int(row.get("safety_stock") or 0),
            "gross_weight_kg": _decimal(row.get("gross_weight_kg")), "length_cm": _decimal(row.get("length_cm")),
            "width_cm": _decimal(row.get("width_cm")), "height_cm": _decimal(row.get("height_cm")), "is_active": True,
        })
    elif resource_type == "channel_mappings":
        store = Store.objects.get(company=company, external_id=str(row.get("store_external_id")))
        sku = SKU.objects.get(company=company, code=str(row.get("internal_sku")))
        entity, _ = ChannelSKU.objects.update_or_create(store=store, external_sku=str(row.get("external_sku")), defaults={
            "sku": sku, "external_product_id": row.get("external_product_id") or "", "fulfillment": row.get("fulfillment") or "fbm",
        })
    elif resource_type == "opening_inventory":
        warehouse = Warehouse.objects.get(company=company, code=str(row.get("warehouse_code")))
        sku = SKU.objects.get(company=company, code=str(row.get("sku")))
        entity, _ = move_inventory(
            company=company, warehouse=warehouse, sku=sku, movement_type=InventoryLedger.Type.OPENING,
            quantity_delta=_decimal(row.get("quantity")), unit_cost=_decimal(row.get("unit_cost")),
            reference_type="import", reference_id=job_id,
            idempotency_key=f"import:{job_id}:{row_number}", actor=actor,
        )
    else:
        raise ValueError(f"暂不支持导入资源: {resource_type}")
    return entity
